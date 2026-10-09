"""AI Resume ATS Score Checker - Streamlit + Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from pypdf import PdfReader

MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
MAX_CHARS = 20000  # keeps prompts small and fast

PROMPT = """You are an expert ATS (Applicant Tracking System) and resume reviewer.
Analyze the resume below{jd_clause} and respond with ONLY valid JSON in this exact schema:

{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {{
    "formatting": <0-100>,
    "keywords": <0-100>,
    "experience": <0-100>,
    "education": <0-100>,
    "skills": <0-100>
  }},
  "strengths": ["<string>", "..."],
  "missing_keywords": ["<string>", "..."],
  "improvements": [
    {{"area": "<string>", "issue": "<string>", "suggestion": "<string>"}}
  ]
}}

Scoring guidance: judge parseability (clear sections, no tables/graphics),
relevant keywords, quantified achievements, action verbs, length, and consistency.
Be honest and specific. Give 3-6 strengths and 5-10 improvements.
{jd_block}
RESUME:
\"\"\"
{resume}
\"\"\"
"""


def extract_text(uploaded_file) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd = (job_description or "").strip()
    return PROMPT.format(
        jd_clause=" against the job description" if jd else "",
        jd_block=f'\nJOB DESCRIPTION:\n"""\n{jd[:MAX_CHARS]}\n"""\n' if jd else "",
        resume=resume_text[:MAX_CHARS],
    )


def parse_response(raw: str) -> dict:
    """Parse the model's JSON, tolerating code fences or stray text."""
    text = (raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if not match:
            raise ValueError("The AI response was not valid JSON. Please try again.")
        data = json.loads(match.group(0))

    def clamp(v):
        try:
            return max(0, min(100, int(round(float(v)))))
        except (TypeError, ValueError):
            return 0

    data["ats_score"] = clamp(data.get("ats_score"))
    data["section_scores"] = {
        k: clamp(v) for k, v in (data.get("section_scores") or {}).items()
    }
    for key in ("strengths", "missing_keywords", "improvements"):
        if not isinstance(data.get(key), list):
            data[key] = []
    data.setdefault("summary", "")
    return data


def get_api_key():
    try:
        key = st.secrets.get("GEMINI_API_KEY")
    except Exception:  # no secrets file locally
        key = None
    return key or os.getenv("GEMINI_API_KEY")


def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict:
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2, response_mime_type="application/json"
        ),
    )
    return parse_response(response.text)


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def main():
    st.set_page_config(page_title="Resume ATS Checker", page_icon="📄", layout="centered")
    st.title("📄 Resume ATS Score Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_api_key()
    if not api_key:
        st.error("Gemini API key not found. Add GEMINI_API_KEY to Streamlit secrets or your environment.")
        st.stop()

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    jd = st.text_area("Job description (optional, for a tailored score)", height=150)

    if st.button("Analyze Resume", type="primary", disabled=uploaded is None):
        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded)
            if len(text) < 50:
                st.error("Couldn't read enough text. If your PDF is a scanned image, upload a text-based PDF or DOCX.")
                st.stop()
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(text, jd, api_key)
        except Exception as e:
            st.error(f"Something went wrong: {e}")
            st.stop()

        score = result["ats_score"]
        st.divider()
        c1, c2 = st.columns([1, 2])
        c1.metric("ATS Score", f"{score}/100", score_label(score))
        c2.write(result["summary"])
        st.progress(score / 100)

        if result["section_scores"]:
            st.subheader("Section scores")
            cols = st.columns(len(result["section_scores"]))
            for col, (name, val) in zip(cols, result["section_scores"].items()):
                col.metric(name.title(), val)

        if result["strengths"]:
            st.subheader("✅ Strengths")
            for s in result["strengths"]:
                st.markdown(f"- {s}")

        if result["missing_keywords"]:
            st.subheader("🔑 Missing keywords")
            st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

        if result["improvements"]:
            st.subheader("🛠️ Suggested improvements")
            for item in result["improvements"]:
                if isinstance(item, dict):
                    with st.expander(item.get("area", "Improvement")):
                        st.markdown(f"**Issue:** {item.get('issue', '')}")
                        st.markdown(f"**Fix:** {item.get('suggestion', '')}")
                else:
                    st.markdown(f"- {item}")

        st.caption("This is an AI estimate, not an official ATS result.")


if __name__ == "__main__":
    main()
